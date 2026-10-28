"""Supplementary tables for the v3 revision package (28 Oct 2026). Every table is built from the ORIGINAL Salmon/tximport
counts (core$sal: DESeq2 ~ NKX3.1 + CM + NKX3.1:CM on all 12 libraries, apeglm shrinkage).

- Supplementary Data 1: DESeq2 results for all tested genes (one sheet per contrast) and normalized counts;
  genes whose Salmon counts come mainly from plasmid reads are marked.
- Supplementary Data 2: GSEA results (salmon_v3/gsea; plasmid-flagged genes removed from the rankings), with the
  expression-trend check (double dagger) and the NKX3-1 removal test, as in v2.
- Supplementary Table S1: every gene labelled in Fig. 1B-D or named in the text (baseMean, normalized counts, unshrunken
  and apeglm-shrunken log2 fold changes), paralogue / readthrough group totals, and the plasmid-flagged genes.

usage: python 68_package_v3_tables.py <out_tables_dir>
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
REV = Path('/mnt/ssd1/Projects/NKX3_wd/results/revision_2026')
CORE = REV / 'core'
V3 = REV / 'salmon_v3'
TAB = V3 / 'gsea' / 'tables'
CIRC = TAB / 'circularity_minus_NKX3-1'
PADJ = 0.05
DAG = '†'
FLAG = pd.read_csv(V3 / 'plasmid_flagged_genes.tsv', sep='\t')
FLAG_IDS = set(FLAG.gene_id)


def readme_sheet(ws, lines, width=140):
    ws.title = 'README'
    for i, (t, bold) in enumerate(lines, start=1):
        c = ws.cell(i, 1, t)
        c.font = Font(bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = width


def write_df(wb, name, df, widths=None, freeze='C2'):
    sh = wb.create_sheet(name)
    sh.append(list(df.columns))
    for row in df.itertuples(index=False):
        sh.append([None if (isinstance(x, float) and pd.isna(x)) else x for x in row])
    for c in sh[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    sh.freeze_panes = freeze
    for j in range(1, len(df.columns) + 1):
        sh.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths and j <= len(widths) else 12)
    return sh


def g3(x):
    return float(f'{x:.3g}') if pd.notna(x) else None


COMMON = [
    ('Quantification: Salmon (v1.10.1, alignment mode on the STAR v2.7.9a transcriptome alignments) within nf-core/rnaseq v3.14.0, '
     'GRCh38 with Ensembl release 99; transcript estimates summarized to genes with tximport (lengthScaledTPM, protein-coding transcripts); '
     'genes with at least 1 count per million in at least 3 libraries (12,965 genes).', False),
    ('Model: one DESeq2 model fitted to all 12 libraries, design ~ NKX3.1 + CM + NKX3.1:CM; main effects from ~ NKX3.1 + CM. '
     'Sign convention: log2 fold change = treatment minus reference (CM minus control; NKX3.1 minus vector control).', False),
    ('Shrunken log2 fold changes: apeglm. Adjusted P: Benjamini-Hochberg. "Adjusted P, test |FC| > 1.5": Wald test against a minimum fold change of 1.5.', False),
    (f'{DAG} Plasmid-derived: reads from the transfected plasmids that STAR v2.7.9a did not align to the genome were written to its transcriptome '
     'alignment (a STAR defect corrected in v2.7.10a) and assigned by Salmon to a few human genes. Genes were flagged when, in a library traced at read '
     'level (one per plasmid group), at least 10 plasmid-derived reads were assigned to the gene and excluding the reads that STAR did not align to the '
     'genome removed at least 50% of its Salmon count (Supplementary Table S1, sheet Plasmid_flagged_genes). Their values do not reflect expression '
     'of the human gene; they are excluded from the gene-set and principal component analyses. NKX3-1 is not flagged: its counts include reads '
     'from the transfected NKX3.1 cDNA, as intended.', False),
]

# ---------------------------------------------------------------- Supplementary Data 1
D1_SHEETS = [('CM_in_vector', 'CM vs control vector cells'), ('NKX_in_ctrl', 'NKX3.1 vs control basal'),
             ('NKX_in_CM', 'NKX3.1 vs control CM'), ('interaction', 'NKX3.1 x CM interaction'),
             ('CM_in_NKX', 'CM vs control NKX3.1 cells'), ('NKX_main', 'NKX3.1 main effect'), ('CM_main', 'CM main effect')]
wb1 = openpyxl.Workbook()
readme_sheet(wb1.active, [('Supplementary Data 1', True),
                          ('Differential expression results (DESeq2) for all 12,965 tested genes.', False)] + COMMON + [
    ('mean_/sd_ columns: DESeq2-normalized counts per group (n = 3).', False),
    ('Sheet "Normalized_counts": DESeq2-normalized counts for every library.', False)])
n_tested = None
for cn, title in D1_SHEETS:
    t = pd.read_csv(CORE / f'deseq2_sal_{cn}.tsv', sep='\t')
    n_tested = len(t)
    df = pd.DataFrame({
        'Ensembl_ID': t.gene_id, 'Gene': t.symbol, 'Biotype': t.biotype,
        f'Plasmid-derived ({DAG})': t.gene_id.map(lambda g: 'yes' if g in FLAG_IDS else ''),
        'baseMean': t.baseMean.round(1), 'log2FC (unshrunken)': t.log2FC_MLE.round(4), 'SE (unshrunken)': t.lfcSE_MLE.round(4),
        'Wald statistic': t.stat.round(4), 'P value': t.pvalue.map(g3), 'Adjusted P (BH)': t.padj.map(g3),
        'log2FC (apeglm-shrunken)': t.log2FC_shrunk.round(4), 'SE (shrunken)': t.lfcSE_shrunk.round(4),
        'Adjusted P, test |FC| > 1.5': t['padj_LFC1.5'].map(g3)})
    for c in [c for c in t.columns if c.startswith(('mean_', 'sd_'))]:
        df[c] = t[c].round(1)
    write_df(wb1, title, df, [17, 14, 16, 10, 10, 11, 10, 10, 10, 10, 11, 10, 11])
nc = pd.read_csv(CORE / 'normalized_counts_sal.tsv', sep='\t')
nc.iloc[:, 2:] = nc.iloc[:, 2:].round(1)
nc.insert(2, f'Plasmid-derived ({DAG})', nc.gene_id.map(lambda g: 'yes' if g in FLAG_IDS else ''))
write_df(wb1, 'Normalized_counts', nc, [17, 14, 10])
wb1.save(OUT / 'Supplementary_Data_1_DESeq2_results.xlsx')
print('Data 1: genes per sheet =', n_tested)

# ---------------------------------------------------------------- Supplementary Data 2
COLLS = [('hallmark', 'Hallmark'), ('gobp', 'GOBP'), ('c2', 'C2')]
CONTRASTS = [('CM_in_vector', 'CMvec'), ('NKX_in_ctrl', 'NKXbas'), ('NKX_in_CM', 'NKXcm'), ('CM_in_NKX', 'CMnkx'),
             ('NKX_main', 'NKXmain'), ('CM_main', 'CMmain'), ('interaction', 'inter')]


def read(q, coll, cn):
    return pd.read_csv(TAB / f'gsea_{q}_{coll}_{cn}.tsv', sep='\t', quoting=3)


def removal_verdict(coll, cn, main):
    ids = main.loc[main.contains_NKX3_1 & (main['p.adjust'] < PADJ), 'ID']
    if ids.empty:
        return pd.DataFrame(columns=['ID', 'NES_without', 'padj_without', 'verdict'])
    wo = pd.read_csv(CIRC / f'gsea_salmon_{coll}_{cn}_minusNKX3-1.tsv', sep='\t', quoting=3)
    d = main.loc[main.ID.isin(ids), ['ID', 'NES', 'p.adjust']].merge(
        wo[['ID', 'NES', 'p.adjust']].rename(columns={'NES': 'NES_without', 'p.adjust': 'padj_without'}), on='ID', how='left')

    def v(r):
        if pd.isna(r.padj_without):
            return 'significant with NKX3-1; too small to test without it'
        if r.padj_without >= PADJ:
            return 'significant only because of NKX3-1'
        return 'robust to removing NKX3-1' if (r.NES > 0) == (r.NES_without > 0) else 'sign reverses without NKX3-1'
    d['verdict'] = d.apply(v, axis=1)
    return d[['ID', 'NES_without', 'padj_without', 'verdict']]


wb2 = openpyxl.Workbook()
readme_sheet(wb2.active, [
    ('Supplementary Data 2', True),
    ('Gene set enrichment analysis (preranked GSEA, clusterProfiler 4.14) for all contrasts.', False),
    ('Genes ranked by the DESeq2 Wald statistic (treatment minus reference) from the model in Supplementary Data 1; '
     f'the plasmid-derived genes ({DAG} in Supplementary Data 1) were removed from every ranking.', False),
    ('Collections: MSigDB Hallmark (H), MSigDB C2 (all curated sets), GO Biological Process. Sets with 15-500 genes.', False),
    ('NES > 0: set higher in the treated group; NES < 0: lower. p.adjust: Benjamini-Hochberg within each collection.', False),
    ('Sheet names: collection_contrast (CMvec = CM vs control in vector cells; NKXbas = NKX3.1 vs control, basal; NKXcm = NKX3.1 vs control under CM; '
     'CMnkx = CM vs control in NKX3.1 cells; NKXmain / CMmain = main effects; inter = NKX3.1 x CM interaction).', False),
    ('', False),
    ('Columns', True),
    ('Leading-edge genes: core enrichment genes reported by GSEA (gene symbols).', False),
    ('NES / p.adjust (trend-adjusted ranking): GSEA repeated after median-centring the Wald statistic within 20 equal-frequency bins of log10 mean expression.', False),
    ('Not robust to expression trend (double dagger in the figures): TRUE for a set with p.adjust < 0.05 that is not significant, or changes sign, with the trend-adjusted ranking. Blank for sets that are not significant.', False),
    ('Contains NKX3-1: the set includes NKX3-1, the transfected gene.', False),
    ('NKX3-1 removal test: for significant sets that contain NKX3-1, GSEA repeated with NKX3-1 removed from the ranking. '
     '"significant only because of NKX3-1", "sign reverses without NKX3-1" and "too small to test without it" mark sets that the figures do not show.', False),
])
summary = []
for coll, cname in COLLS:
    for cn, cshort in CONTRASTS:
        main = read('salmon', coll, cn)
        adj = read('salmonadj', coll, cn)[['ID', 'NES', 'p.adjust']].rename(columns={'NES': 'NES_adj', 'p.adjust': 'padj_adj'})
        d = main.merge(adj, on='ID', how='left')
        sig = d['p.adjust'] < PADJ
        nonrob = d.padj_adj.isna() | (d.padj_adj >= PADJ) | ((d.NES_adj > 0) != (d.NES > 0))
        d['ddag'] = pd.Series([bool(x) if s else None for x, s in zip(nonrob, sig)], dtype=object)
        d = d.merge(removal_verdict(coll, cn, main), on='ID', how='left').sort_values('p.adjust', kind='stable')
        out = pd.DataFrame({
            'ID': d.ID, 'Description': d.Description, 'setSize': d.setSize, 'NES': d.NES.round(3),
            'pvalue': d.pvalue.map(g3), 'p.adjust': d['p.adjust'].map(g3),
            'NES (trend-adjusted ranking)': d.NES_adj.round(3), 'p.adjust (trend-adjusted ranking)': d.padj_adj.map(g3),
            'Not robust to expression trend': d.ddag, 'Contains NKX3-1': d.contains_NKX3_1.map({True: 'yes', False: ''}),
            'NKX3-1 removal test': d.verdict.fillna(''), 'Leading-edge genes': d.core_enrichment_symbols})
        write_df(wb2, f'{cname}_{cshort}', out, [44, 44, 8, 8, 10, 10, 12, 12, 12, 10, 34, 80])
        sig = sig.reindex(d.index)
        summary.append((cname, cshort, len(out), int((d['p.adjust'] < PADJ).sum()), int(((d['p.adjust'] < PADJ) & (d.NES > 0)).sum()),
                        int(d.ddag.fillna(False).astype(bool).sum())))
wb2.save(OUT / 'Supplementary_Data_2_GSEA_results.xlsx')
print(pd.DataFrame(summary, columns=['coll', 'contrast', 'n_sets', 'n_sig', 'n_sig_up', 'n_sig_not_trend_robust']).to_string())

# ---------------------------------------------------------------- Supplementary Table S1
ng = pd.read_csv(V3 / 'named_genes' / 'named_genes.tsv', sep='\t')
nn = pd.read_csv(V3 / 'named_genes' / 'named_genes_normalized_counts.tsv', sep='\t')
fam = pd.read_csv(V3 / 'named_genes' / 'family_totals.tsv', sep='\t')
CT = [('CM_in_vector', 'CM vs control, vector cells (Fig. 1B)'), ('NKX_in_ctrl', 'NKX3.1 vs vector, basal (Fig. 1C)'),
      ('NKX_in_CM', 'NKX3.1 vs vector, CM (Fig. 1D)'), ('interaction', 'NKX3.1 x CM interaction')]
s1 = pd.DataFrame({'Gene': ng.symbol, 'Ensembl_ID': ng.gene_id, 'Biotype': ng.biotype,
                   'Labelled in Fig. 1 panel(s)': ng.labelled_in.fillna(''),
                   'Reported in the submitted manuscript (text or Fig. 1E-G)': ng.named_in_submitted_text.fillna(''),
                   f'Plasmid-derived ({DAG})': ng.plasmid_derived.fillna(''),
                   'Paralogue / readthrough group (see sheet Paralogue_readthrough_groups)': ng.family_or_locus.fillna(''),
                   'baseMean': ng.baseMean.round(1)})
for cn, lab in CT:
    s1[f'{lab}: log2FC (unshrunken)'] = ng[f'{cn}_log2FC_MLE'].round(3)
    s1[f'{lab}: SE (unshrunken)'] = ng[f'{cn}_SE_MLE'].round(3)
    s1[f'{lab}: log2FC (apeglm-shrunken)'] = ng[f'{cn}_log2FC_apeglm'].round(3)
    s1[f'{lab}: SE (shrunken)'] = ng[f'{cn}_SE_apeglm'].round(3)
    s1[f'{lab}: adjusted P'] = ng[f'{cn}_padj'].map(g3)
for g in ['Control-basal', 'Control-inflammation', 'NKX3.1-basal', 'NKX3.1-inflammation']:
    s1[f'{g}: mean normalized count'] = ng[f'mean_{g}'].round(1)
    s1[f'{g}: SD'] = ng[f'sd_{g}'].round(1)
nn2 = nn.copy()
nn2.iloc[:, 2:] = nn2.iloc[:, 2:].round(1)
nn2.insert(2, f'Plasmid-derived ({DAG})', nn2.gene_id.map(lambda g: 'yes' if g in FLAG_IDS else ''))
famcols = [c for c in fam.columns if c not in ('family', 'contrast', 'baseMean', 'log2FC_MLE', 'SE_MLE', 'log2FC_apeglm', 'SE_apeglm', 'padj', 'members')]
fam_out = fam.rename(columns={'family': 'Group', 'members': 'Members (counts summed per library)', 'contrast': 'Contrast',
                              'log2FC_MLE': 'log2FC (unshrunken)', 'SE_MLE': 'SE (unshrunken)', 'log2FC_apeglm': 'log2FC (apeglm-shrunken)',
                              'SE_apeglm': 'SE (shrunken)', 'padj': 'Adjusted P'})
fam_out['Contrast'] = fam_out['Contrast'].map(dict(CT))
fam_out = fam_out[['Group', 'Members (counts summed per library)', 'Contrast', 'baseMean', 'log2FC (unshrunken)', 'SE (unshrunken)',
                   'log2FC (apeglm-shrunken)', 'SE (shrunken)', 'Adjusted P'] + famcols]
for c in ['baseMean', 'log2FC (unshrunken)', 'SE (unshrunken)', 'log2FC (apeglm-shrunken)', 'SE (shrunken)']:
    fam_out[c] = fam_out[c].round(3)
fam_out['Adjusted P'] = fam_out['Adjusted P'].map(g3)
fl = FLAG.rename(columns={'library': 'Traced library (plasmid group)', 'gene_id': 'Ensembl_ID', 'gene_name': 'Gene',
                          'nfcore_original': 'Salmon count (nf-core run)',
                          'salmon_without_unmapped': 'Salmon count without reads STAR did not align to the genome',
                          'plasmid_reads': 'Plasmid-derived reads assigned to the gene', 'frac_count_removed': 'Fraction of Salmon count removed'})
fl['Passes the expression filter (in Supplementary Data 1)'] = fl.Ensembl_ID.map(lambda g: 'yes' if g in set(nc.gene_id) else 'no')

wb3 = openpyxl.Workbook()
readme_sheet(wb3.active, [
    ('Supplementary Table S1', True),
    ('Expression values for every gene labelled in Fig. 1B-D (the 20 genes with the largest |z| per panel), named in the revised text, or reported as an RNA-seq result in the submitted manuscript (text and Fig. 1E-G).', False)] + COMMON + [
    ('', False),
    ('Sheets', True),
    ('Named_genes: baseMean, unshrunken and apeglm-shrunken log2 fold changes with standard errors and adjusted P values for the contrasts of '
     'Fig. 1B-D and the NKX3.1 x CM interaction, and DESeq2-normalized counts per group (mean and SD, n = 3).', False),
    ('Normalized_counts: DESeq2-normalized counts for every library.', False),
    ('Paralogue_readthrough_groups: Salmon distributes reads that are compatible with several transcripts between them (expectation-maximization), '
     'so the split between near-identical paralogues or between a gene and a readthrough transcript is uncertain. For the groups named by '
     'Reviewer 1 (SLX1A/SLX1B, RGPD family, COMMD3-BMI1, AC131160.1), the GH/CSH cluster and ASDURF/ASNSD1, member counts were summed per '
     'library and tested as one feature in the same model (size factors of the gene-level fit). Columns after "Adjusted P": normalized group totals per library.', False),
    ('Plasmid_flagged_genes: the read-tracing evidence behind the plasmid flag; genes that do not pass the expression filter are listed for completeness.', False),
    ('Genes are listed by panel label first, then by baseMean.', False),
])
write_df(wb3, 'Named_genes', s1, [12, 17, 15, 10, 12, 10, 30, 10])
write_df(wb3, 'Normalized_counts', nn2, [17, 12, 10])
write_df(wb3, 'Paralogue_readthrough_groups', fam_out, [40, 40, 34, 10, 10, 10, 10, 10, 10], freeze='B2')
write_df(wb3, 'Plasmid_flagged_genes', fl, [30, 17, 13, 12, 14, 14, 12, 14], freeze='C2')
wb3.save(OUT / 'Supplementary_Table_S1_named_genes.xlsx')
print('S1:', len(s1), 'genes;', len(fam_out), 'group rows;', len(fl), 'flagged genes')
