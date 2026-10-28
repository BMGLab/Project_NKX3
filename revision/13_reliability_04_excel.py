#!/usr/bin/env python
"""Reliability task, step 4: write the reviewer-requested supplementary table (Excel + TSV).
Input: outputs of 13_reliability_02_bam_audit.py and 13_reliability_03_table.R.
Run: /home/biolab/miniconda3/envs/nkx3_py/bin/python -I 13_reliability_04_excel.py
"""
import json
import numpy as np
import pandas as pd
import warnings
warnings.simplefilter('ignore', pd.errors.PerformanceWarning)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = '/mnt/ssd1/Projects/NKX3_wd/results/revision_2026/reliability'
XLSX = f'{OUT}/SuppTable_named_genes_quantification.xlsx'
GRP = ['Control-basal', 'Control-inflammation', 'NKX3.1-basal', 'NKX3.1-inflammation']
CTX = [('CM_in_vector', 'CM vs control (vector cells)'), ('NKX_in_ctrl', 'NKX3.1 vs control (basal)'),
       ('NKX_in_CM', 'NKX3.1 vs control (inflammation)'), ('interaction', 'Interaction (NKX3.1 x CM)')]
MAINS = [('NKX_main', 'NKX3.1 main effect (additive model)'), ('CM_main', 'CM main effect (additive model)')]
STATS = [('log2FC_MLE', 'log2FC (MLE)'), ('log2FC_apeglm', 'log2FC (apeglm)'), ('lfcSE_apeglm', 'SE (apeglm)'),
         ('Wald_stat', 'Wald stat'), ('padj', 'padj (BH)'), ('padj_LFC1.5', 'padj, |FC|>1.5 test')]
QUANT = [('corrected', 'CORRECTED quantification: STAR genome alignments + featureCounts (unique fragments)'),
         ('original', 'ORIGINAL quantification (submitted manuscript): STAR -> Salmon, tximport lengthScaledTPM')]

full = pd.read_csv(f'{OUT}/named_genes_reliability_full.tsv', sep='\t')
full['tier_code'] = full.tier.str[0]
full = full.sort_values(['tier_code', 'symbol'], key=lambda s: s if s.name == 'tier_code' else s.str.upper()).reset_index(drop=True)
# keep the curated order within tier A (manuscript order) for readability
cur = pd.read_csv(f'{OUT}/named_genes_curated.tsv', sep='\t')
full['ord'] = full.symbol.map({s: i for i, s in enumerate(cur.symbol)})
full = full.sort_values(['tier_code', 'ord']).reset_index(drop=True)
evs = pd.read_csv(f'{OUT}/read_evidence_per_sample.tsv', sep='\t')
fam = pd.read_csv(f'{OUT}/family_totals_tests.tsv', sep='\t')
txc = pd.read_csv(f'{OUT}/salmon_transcript_counts_audited_genes.tsv', sep='\t')
stats = pd.read_csv(f'{OUT}/star_mapping_stats.tsv', sep='\t')

FILL = {'id': 'DDDDDD', 'ev': 'FCE4D6', 'corrected': 'E2EFDA', 'original': 'FFF2CC', 'published': 'D9E1F2'}
thin = Side(style='thin', color='999999')
HFONT = Font(name='Arial', bold=True, size=9)
BFONT = Font(name='Arial', size=9)
NUMFMT = {'count': '#,##0', 'norm': '#,##0.0', 'lfc': '0.00', 'se': '0.00', 'stat': '0.00', 'p': '0.00E+00',
          'frac': '0.00', 'text': '@', 'fold': '0.000'}


def write_block_sheet(ws, df, blocks, freeze='C3'):
    """blocks: list of (block_title, fill_key, [(col_name_in_df, header, fmt), ...])"""
    c = 1
    for title, fk, cols in blocks:
        c0 = c
        for col, hdr, fmt in cols:
            h = ws.cell(row=2, column=c, value=hdr)
            h.font = HFONT; h.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center')
            h.fill = PatternFill('solid', fgColor=FILL[fk]); h.border = Border(bottom=thin, top=thin, left=thin, right=thin)
            for i, v in enumerate(df[col].tolist() if col in df.columns else [None] * len(df)):
                if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
                    v = None
                if isinstance(v, (np.integer,)):
                    v = int(v)
                if isinstance(v, (np.floating,)):
                    v = float(v)
                if isinstance(v, (np.bool_, bool)):
                    v = 'yes' if v else 'no'
                cell = ws.cell(row=3 + i, column=c, value=v)
                cell.font = BFONT
                if fmt != 'text':
                    cell.number_format = NUMFMT[fmt]
                else:
                    cell.alignment = Alignment(wrap_text=False, vertical='top')
            width = {'text': 18, 'count': 9, 'norm': 9, 'lfc': 8, 'se': 7, 'stat': 7, 'p': 9, 'frac': 7, 'fold': 9}[fmt]
            if hdr in ('Symbol',):
                width = 13
            if hdr in ('Read-evidence class',):
                width = 26
            if hdr in ('Reason', 'Where named in manuscript', 'Corrected-analysis result for the reported contrast(s)',
                       'Family / locus total test (genome fragments)'):
                width = 60
            ws.column_dimensions[get_column_letter(c)].width = width
            c += 1
        if c - 1 >= c0:
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c - 1)
            t = ws.cell(row=1, column=c0, value=title)
            t.font = Font(name='Arial', bold=True, size=10); t.fill = PatternFill('solid', fgColor=FILL[fk])
            t.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 54
    ws.freeze_panes = freeze


def quant_cols(q, which):
    cols = [(f'{q}_prefilter', 'Pre-filter status (>=10 reads in >=3 samples)', 'text'),
            (f'{q}_baseMean', 'baseMean (mean normalised count, 12 samples)', 'norm')]
    for g in GRP:
        cols += [(f'{q}_normcount_mean_{g}', f'{g}: normalised count, mean', 'norm'),
                 (f'{q}_normcount_sd_{g}', f'{g}: normalised count, SD', 'norm')]
    for g in GRP:
        cols.append((f'{q}_rawcount_mean_{g}', f'{g}: raw count, mean', 'norm'))
    for key, lab in which:
        for st, sl in STATS:
            fmt = {'log2FC_MLE': 'lfc', 'log2FC_apeglm': 'lfc', 'lfcSE_apeglm': 'se', 'Wald_stat': 'stat'}.get(st, 'p')
            cols.append((f'{q}_{key}_{st}', f'{lab}: {sl}', fmt))
    return cols


ID_COLS = [('symbol', 'Symbol', 'text'), ('gene_id', 'Ensembl gene ID (Ensembl 99)', 'text'), ('biotype', 'Biotype', 'text'),
           ('locus', 'Locus (GRCh38)', 'text'), ('tier', 'Tier', 'text'), ('where_named', 'Where named in manuscript', 'text'),
           ('reported_claims', 'Change reported in manuscript (contrast:direction)', 'text')]
EV_COLS = [('read_evidence_class', 'Read-evidence class', 'text'), ('read_evidence_reason', 'Reason', 'text'),
           ('corrected_reported_contrast_result', 'Corrected-analysis result for the reported contrast(s)', 'text'),
           ('family_level_test', 'Family / locus total test (genome fragments)', 'text')]
PUB_COLS = []
for key, lab in CTX[:3]:
    PUB_COLS += [(f'published_{key}_log2FC', f'{lab}: published log2FC (sign = treatment - reference)', 'lfc'),
                 (f'published_{key}_fold', f'{lab}: published fold change', 'fold'),
                 (f'published_{key}_padj', f'{lab}: published padj', 'p')]

wb = Workbook()
# ------------------------------------------------------------------ README
ws = wb.active; ws.title = 'README'
n_cls = full.groupby(['tier_code', 'read_evidence_class']).size().unstack(fill_value=0)
readme = [
    ('Supplementary Table SX. Expression evidence for every gene named in the manuscript text and figures', 'title'),
    ('', ''),
    ('Design', 'h'),
    ('DU145 cells, 4 groups x 3 biological replicates: Control-basal (empty vector, no CM), Control-inflammation (empty vector + macrophage-conditioned medium, CM), NKX3.1-basal (NKX3.1 plasmid, no CM), NKX3.1-inflammation (NKX3.1 plasmid + CM). Stranded (reverse) paired-end mRNA-seq.', ''),
    ('Sign convention', 'h'),
    ('All log2 fold changes are treatment minus reference: CM vs control and NKX3.1 vs empty vector. (The submitted Fig 1B, 2A, 4A and Supp Fig 3B used control minus CM; the "published" columns here have been converted to treatment minus reference.)', ''),
    ('Two quantifications', 'h'),
    ('CORRECTED (used for the revision): STAR 2.7.9a genome alignments from nf-core/rnaseq 3.14.0 (GRCh38 primary assembly, Ensembl 99 annotation), counted with featureCounts 2.1.1 (-p --countReadPairs -s 2 -t exon -g gene_id). Each read pair (fragment) is counted once; only uniquely aligned fragments (NH:i:1) are counted; fragments overlapping exons of more than one gene are left unassigned (featureCounts default, no -O); multi-mapped fragments (NH>1) are not counted (no -M); PCR duplicates are retained. 60,676 annotated genes.', ''),
    ('ORIGINAL (submitted manuscript): Salmon 1.10.1 in alignment mode on the STAR transcriptome BAM (nf-core/rnaseq 3.14.0), imported with tximport (countsFromAbundance = "lengthScaledTPM") with a protein-coding-only tx2gene, then CPM >= 1 in >= 3 samples (12,965 genes). In this mode Salmon distributes multi-mapping fragments among paralogues and among overlapping/readthrough transcripts by expectation-maximisation; no gene-specific handling was applied.', ''),
    ('Statistical model (both quantifications)', 'h'),
    ('DESeq2 1.46 on all 12 samples with design ~ NKX + CM + NKX:CM. Per-context contrasts (CM vs control in vector cells; NKX3.1 vs control under basal and under inflammatory conditions) and the interaction are model coefficients after re-levelling. NKX3.1 and CM main effects come from the additive model ~ NKX + CM (sheet Main_effects). Pre-filter: >= 10 reads in >= 3 samples (the size of one group). Wald tests, Benjamini-Hochberg adjustment over all tested genes. "padj, |FC|>1.5 test" tests H0: |log2FC| <= log2(1.5) (lfcThreshold = log2(1.5), altHypothesis = "greaterAbs").', ''),
    ('LFC shrinkage: apeglm (lfcShrink type = "apeglm") for every contrast. apeglm learns its prior from all genes; because almost no gene changes in this experiment, the prior is narrow and shrunken log2FCs are ~0 except where the data are very strong (e.g. NKX3-1). MLE estimates are shown alongside. Applied to the original counts, apeglm does not remove the artefactual fold changes (e.g. BCL11B +9.5, CSH2 -8.1), because the problem lies in the counts, not in the estimator.', ''),
    ('Normalised counts = raw counts / DESeq2 median-of-ratios size factors of the full model. They are shown for every gene, including genes that did not pass the pre-filter (DESeq2 statistics blank).', ''),
    ('As published', 'h'),
    ('Values behind Fig 1E-G / Supp Fig 2A-C: separate pairwise DESeq2 fits on the original counts (rna_downstreams_out_v15/dgea_*), converted to treatment minus reference. The figure bars showed these fold changes without error bars.', ''),
    ('Read-evidence audit (sheets Read_evidence, Per_sample_evidence, Family_totals)', 'h'),
    ('For each gene and sample, fragments (read pairs, counted once per read name) were counted directly from the coordinate-sorted STAR genome BAMs (pysam, indexed region queries), strand-aware for the reverse-stranded library (read 1 antisense to the transcript; Salmon libType ISR). Secondary alignments of multi-mapped fragments are included. Categories: sense fragments with any aligned block overlapping the union of the gene\'s Ensembl 99 exons ("genome, overlapping"); sense fragments whose aligned blocks all lie within those exons (+/- 3 nt) with the mate inside the gene ("genome, exon-contained", the transcript-compatible fragments that Salmon could assign); unique (NH:i:1) vs multi-mapped (NH>1); duplicate-flagged (0x400); shared (also overlapping exons of another gene on the same strand, e.g. readthrough or nested genes); antisense exonic; sense intronic. Compared with the nf-core Salmon gene count (salmon.merged.gene_counts.tsv, sum of transcript NumReads) and featureCounts.', ''),
    ('Classes (applied in this order): (1) "low expression": every source < 10 fragments per sample on average in every group. (2) "not supported by read evidence": in at least one group Salmon assigns >= 10 fragments and > 1.5x (and >= 10 more than) the number of sense fragments overlapping the gene\'s exons in the genome alignments, counting unique and multi-mapped alignments, i.e. Salmon reports fragments that do not exist at the locus; genes with Salmon/genome disagreement or opposite directions not explained by mappability also fall here. (3) "mappability-limited": >= 10 fragments at the locus but >= 50% of them multi-mapped (NH>1) or shared with exons of another same-strand gene, so the gene-level estimate is not identifiable; family / locus totals (distinct fragments across members) are reported and tested instead. (4) "supported": Salmon within 1.5-fold of the exon-contained genome fragments in every group with >= 10 fragments, and the reported direction of change agrees between Salmon and genome counts. "Supported" refers to the quantification only: whether the reported change is statistically reproduced is given in the column "Corrected-analysis result for the reported contrast(s)".', ''),
    ('Family / locus totals were tested with DESeq2 (same design and size factors; rows appended to the full count matrix) and with a likelihood-ratio test of any group difference (~group vs ~1).', ''),
    ('Tiers', 'h'),
    ('A: gene reported as a result or marker (Abstract, Results, Discussion, Fig 1E-G, Supp Fig 2A-C; CD44/PROM1/CDH2 are the flow-cytometry and western-blot markers). B: other genes named in the text (context) and two epithelial markers named by Reviewer 1. C: genes visible only as figure labels (top-20 |z| labels of Fig 1B-D / Supp Fig 2D-F; module members in Fig 3D-E).', ''),
    ('Summary of classes', 'h'),
]
for tcode in n_cls.index:
    readme.append((f'Tier {tcode}: ' + ', '.join(f'{k} {v}' for k, v in n_cls.loc[tcode].items() if v), ''))
readme += [('Sheets', 'h'),
           ('Named_genes: one row per gene, both quantifications side by side (per-context contrasts + interaction) and the published values. Main_effects: NKX3.1 and CM main effects (additive model). Read_evidence: per-gene audit summary (group means of fragments by source). Family_totals: paralogue families (SLX1A+SLX1B; RGPD1-RGPD8; GH1+GH2+CSH1+CSH2+CSHL1; COMMD3+BMI1+COMMD3-BMI1) and locus families (named gene + same-strand exon-overlapping genes), with tests. Per_sample_evidence: every count per gene and sample. Salmon_transcripts: nf-core Salmon NumReads per transcript for the audited genes. Mapping_stats: STAR read-pair counts per sample.', ''),
           ('Scripts: /mnt/ssd1/Projects/NKX3_wd/scripts/revision/00_core_deseq2.R and 13_reliability_0[0-5]_* (run all with 13_reliability_00_run_all.sh).', '')]
ws.column_dimensions['A'].width = 160
for i, (txt, kind) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    c.font = Font(name='Arial', size=12 if kind == 'title' else 10, bold=kind in ('title', 'h'))

# ------------------------------------------------------------------ Named_genes
ws = wb.create_sheet('Named_genes')
blocks = [('Gene', 'id', ID_COLS), ('Read-evidence audit and corrected result', 'ev', EV_COLS)]
for q, lab in QUANT:
    blocks.append((lab, q, quant_cols(q, CTX)))
blocks.append(('AS PUBLISHED (original pairwise DESeq2 behind Fig 1E-G; unshrunken)', 'published', PUB_COLS))
write_block_sheet(ws, full, blocks)
ws.auto_filter.ref = f'A2:{get_column_letter(ws.max_column)}{ws.max_row}'

# ------------------------------------------------------------------ Main_effects
ws = wb.create_sheet('Main_effects')
blocks = [('Gene', 'id', ID_COLS[:5] + [('read_evidence_class', 'Read-evidence class', 'text')])]
for q, lab in QUANT:
    cols = [(f'{q}_prefilter', 'Pre-filter status', 'text')]
    for key, lab2 in MAINS:
        for st, sl in STATS:
            fmt = {'log2FC_MLE': 'lfc', 'log2FC_apeglm': 'lfc', 'lfcSE_apeglm': 'se', 'Wald_stat': 'stat'}.get(st, 'p')
            cols.append((f'{q}_{key}_{st}', f'{lab2}: {sl}', fmt))
    blocks.append((lab, q, cols))
write_block_sheet(ws, full, blocks)

# ------------------------------------------------------------------ Read_evidence
ws = wb.create_sheet('Read_evidence')
cols = []
for src, lab in [('salmon', 'Salmon (nf-core NumReads)'), ('genome_overlap', 'Genome, overlapping exons (unique+multi)'),
                 ('genome_contained', 'Genome, exon-contained (unique+multi)'), ('genome_unique', 'Genome, unique (NH=1)'),
                 ('featurecounts', 'featureCounts (unique, gene-specific)')]:
    for g in GRP:
        cols.append((f'{src}_mean_{g}', f'{lab}: {g} mean', 'norm'))
frac = [('multi_mapped_fraction', 'Fraction multi-mapped (NH>1)', 'frac'), ('shared_fraction', 'Fraction shared with another same-strand gene', 'frac'),
        ('duplicate_fraction', 'Fraction duplicate-flagged', 'frac'), ('salmon_to_genome_ratio_min', 'Salmon / exon-contained, min over groups', 'frac'),
        ('salmon_to_genome_ratio_max', 'Salmon / exon-contained, max over groups', 'frac'),
        ('antisense_exonic_total', 'Antisense exonic fragments (12 samples)', 'count'), ('sense_intronic_total', 'Sense intronic fragments (12 samples)', 'count'),
        ('reported_dir_salmon', 'Reported contrast: log2 ratio, Salmon (normalised means)', 'text'),
        ('reported_dir_genome', 'Reported contrast: log2 ratio, genome exon-contained', 'text'),
        ('same_strand_exon_overlap_genes', 'Same-strand genes sharing exonic sequence', 'text'), ('families', 'Family / locus groups', 'text')]
blocks = [('Gene', 'id', [ID_COLS[0], ID_COLS[1], ID_COLS[4], ('read_evidence_class', 'Read-evidence class', 'text'),
                          ('read_evidence_reason', 'Reason', 'text')]),
          ('Mean fragments per sample in each group', 'ev', cols), ('Locus properties', 'original', frac),
          ('Family / locus total', 'corrected', [('family_level_test', 'Family / locus total test (genome fragments)', 'text')])]
write_block_sheet(ws, full, blocks)

# ------------------------------------------------------------------ Family_totals
ws = wb.create_sheet('Family_totals')
fcols = [('family', 'Family / locus', 'text'), ('members', 'Members', 'text')]
mcols = []
for src, lab in [('genome', 'Genome distinct fragments'), ('genome_unique', 'Genome unique'), ('salmon', 'Salmon sum'),
                 ('featurecounts', 'featureCounts sum')]:
    for g in GRP:
        mcols.append((f'{src}_mean_{g}', f'{lab}: {g} mean', 'norm'))
tcols = [('genome_LRT_padj', 'Genome: LRT any group difference, padj', 'p'), ('salmon_LRT_padj', 'Salmon: LRT any group difference, padj', 'p')]
for key, lab in CTX:
    for src in ('genome', 'salmon'):
        tcols += [(f'{src}_{key}_log2FC_MLE', f'{src.capitalize()}: {lab}: log2FC (MLE)', 'lfc'),
                  (f'{src}_{key}_padj', f'{src.capitalize()}: {lab}: padj', 'p')]
fam_sorted = pd.concat([fam[~fam.family.str.startswith('locus:')], fam[fam.family.str.startswith('locus:')]])
write_block_sheet(ws, fam_sorted, [('Family', 'id', fcols), ('Mean fragments per sample', 'ev', mcols), ('Tests (DESeq2)', 'corrected', tcols)])

# ------------------------------------------------------------------ Per_sample_evidence
ws = wb.create_sheet('Per_sample_evidence')
pcols = [('symbol', 'Symbol', 'text'), ('gene_id', 'Ensembl gene ID', 'text'), ('sample_id', 'Sample', 'text'), ('label', 'Group', 'text'),
         ('salmon_numreads', 'Salmon NumReads (nf-core gene count)', 'norm'), ('featurecounts_unique', 'featureCounts (unique)', 'count'),
         ('sense_exonic_all', 'Sense, overlapping exons: all', 'count'), ('sense_contained_all', 'Sense, exon-contained: all', 'count'),
         ('sense_exonic_unique', 'Sense exonic: unique (NH=1)', 'count'), ('sense_exonic_multi', 'Sense exonic: multi-mapped (NH>1)', 'count'),
         ('sense_exonic_multi_primary_here', 'Sense exonic multi-mapped with primary alignment here', 'count'),
         ('sense_exonic_shared', 'Sense exonic: shared with another same-strand gene', 'count'),
         ('sense_exonic_unique_specific', 'Sense exonic: unique and gene-specific', 'count'),
         ('sense_contained_unique_specific', 'Sense exon-contained: unique and gene-specific', 'count'),
         ('sense_exonic_dup', 'Sense exonic: duplicate-flagged', 'count'), ('sense_exonic_unique_nodup', 'Sense exonic: unique, not duplicate', 'count'),
         ('antisense_exonic_all', 'Antisense exonic: all', 'count'), ('antisense_exonic_unique', 'Antisense exonic: unique', 'count'),
         ('sense_intronic_all', 'Sense intronic: all', 'count'), ('sense_intronic_unique', 'Sense intronic: unique', 'count'),
         ('sense_exonic_mean_NM', 'Mean edit distance (NM) of sense exonic fragments', 'frac')]
order_s = {s: i for i, s in enumerate(cur.symbol)}
evs['o'] = evs.symbol.map(order_s).fillna(999)
evs = evs.sort_values(['o', 'symbol', 'group', 'sample_id'], key=lambda s: s.map({'vector_ctrl': 0, 'vector_CM': 1, 'NKX3.1_ctrl': 2, 'NKX3.1_CM': 3}) if s.name == 'group' else s)
write_block_sheet(ws, evs, [('Per gene x sample fragment counts (genome BAM audit) and quantifications', 'ev', pcols)], freeze='B3')

# ------------------------------------------------------------------ Salmon_transcripts
ws = wb.create_sheet('Salmon_transcripts')
SIDS = list(stats.sample_id)
lab_of = dict(zip(evs.sample_id, evs.label))
txc = txc[txc[SIDS].sum(axis=1) >= 1].copy()
txc['o'] = txc.symbol.map(order_s).fillna(999)
txc = txc.sort_values(['o', 'symbol', 'tx'])
tcols = [('symbol', 'Gene', 'text'), ('tx', 'Transcript ID', 'text'), ('transcript_name', 'Transcript name', 'text'),
         ('transcript_biotype', 'Transcript biotype', 'text'), ('transcript_length', 'Length (nt)', 'count')] + \
        [(s, f'{s} ({lab_of[s]})', 'norm') for s in SIDS]
write_block_sheet(ws, txc, [('nf-core Salmon NumReads per transcript (salmon.merged.transcript_counts.tsv), audited genes', 'ev', tcols)], freeze='C3')

# ------------------------------------------------------------------ Mapping_stats
ws = wb.create_sheet('Mapping_stats')
stats['label'] = stats.sample_id.map(lab_of)
write_block_sheet(ws, stats, [('STAR (nf-core) read pairs per sample', 'id',
                               [('sample_id', 'Sample', 'text'), ('label', 'Group', 'text'), ('input_pairs', 'Input pairs', 'count'),
                                ('unique_pairs', 'Uniquely mapped pairs', 'count'), ('multi_pairs', 'Multi-mapped pairs (2-20 loci)', 'count'),
                                ('toomany_pairs', 'Pairs mapped to >20 loci', 'count')])], freeze='B3')
wb.save(XLSX)

# ------------------------------------------------------------------ flat TSVs
def flat(df, blocks):
    cols, names = [], []
    for title, fk, cs in blocks:
        for col, hdr, fmt in cs:
            if col in df.columns:
                cols.append(col); names.append(col)
    return df[cols]


main_blocks = [('Gene', 'id', ID_COLS), ('ev', 'ev', EV_COLS)] + [(lab, q, quant_cols(q, CTX)) for q, lab in QUANT] + [('pub', 'published', PUB_COLS)]
flat(full, main_blocks).to_csv(f'{OUT}/SuppTable_named_genes_quantification.tsv', sep='\t', index=False)
me_blocks = [('Gene', 'id', ID_COLS[:5] + [('read_evidence_class', '', 'text')])] + \
            [(lab, q, [(f'{q}_prefilter', '', 'text')] + [(f'{q}_{k}_{st}', '', '') for k, _ in MAINS for st, _ in STATS]) for q, lab in QUANT]
flat(full, me_blocks).to_csv(f'{OUT}/SuppTable_named_genes_main_effects.tsv', sep='\t', index=False)
print('wrote', XLSX)


def _write_versions():
    import sys, importlib, platform
    with open(f'{OUT}/sessionInfo_13_reliability_04_excel_python.txt', 'w') as fh:
        fh.write(f'python {sys.version}\nplatform {platform.platform()}\n')
        for m in ['openpyxl', 'numpy', 'pandas']:
            fh.write(f'{m} {importlib.import_module(m).__version__}\n')


_write_versions()
