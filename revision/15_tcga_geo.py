#!/usr/bin/env python
"""15_tcga_geo.py -- time-boxed GEO search for public NKX3.1 perturbation expression datasets in prostate
cell lines (Reviewer 2 mentions PC-3 +/- NKX3.1; Reviewer 1 major 5 asks for independent data).

Queries NCBI E-utilities (db=gds, series records), stores the raw JSON, and writes every hit with
title / summary / platform / sample count, plus keyword flags used for manual curation.

Run: /home/biolab/miniconda3/envs/nkx3_py/bin/python -I /mnt/ssd1/Projects/NKX3_wd/scripts/revision/15_tcga_geo.py
"""
import json, time, urllib.parse, urllib.request, re, csv, datetime, os

OUT = "/mnt/ssd1/Projects/NKX3_wd/results/revision_2026/tcga/tables"
WORK = "/mnt/ssd1/Projects/NKX3_wd/work/tcga/geo"
os.makedirs(OUT, exist_ok=True); os.makedirs(WORK, exist_ok=True)
EU = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/"
TOOL = "nkx3_revision_geo_search"

QUERIES = {
    "nkx31_human_gse": '(NKX3-1[All Fields] OR NKX3.1[All Fields] OR NKX31[All Fields]) AND "Homo sapiens"[Organism] AND gse[Entry Type]',
    "nkx31_prostate_any_org": '(NKX3-1[All Fields] OR NKX3.1[All Fields]) AND prostate[All Fields] AND gse[Entry Type]',
    "nkx31_pc3": '(NKX3-1[All Fields] OR NKX3.1[All Fields]) AND (PC-3[All Fields] OR PC3[All Fields]) AND gse[Entry Type]',
    "nkx31_du145": '(NKX3-1[All Fields] OR NKX3.1[All Fields]) AND DU145[All Fields] AND gse[Entry Type]',
    "nkx31_lncap": '(NKX3-1[All Fields] OR NKX3.1[All Fields]) AND LNCaP[All Fields] AND gse[Entry Type]',
}


def get(url, params):
    params = dict(params, tool=TOOL, retmode="json")
    q = url + "?" + urllib.parse.urlencode(params)
    for attempt in range(4):
        try:
            with urllib.request.urlopen(q, timeout=60) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception as e:  # transient network errors
            time.sleep(1.5 * (attempt + 1))
            last = e
    raise last


def main():
    accessed = datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
    hits = {}
    raw = {"accessed_utc": accessed, "queries": {}, "summaries": {}}
    for name, term in QUERIES.items():
        r = get(EU + "esearch.fcgi", {"db": "gds", "term": term, "retmax": 500})
        ids = r["esearchresult"].get("idlist", [])
        raw["queries"][name] = {"term": term, "count": r["esearchresult"].get("count"), "ids": ids}
        for i in ids:
            hits.setdefault(i, set()).add(name)
        time.sleep(0.4)
    uids = sorted(hits)
    for k in range(0, len(uids), 100):
        r = get(EU + "esummary.fcgi", {"db": "gds", "id": ",".join(uids[k:k + 100])})
        for u in r["result"].get("uids", []):
            raw["summaries"][u] = r["result"][u]
        time.sleep(0.4)
    with open(os.path.join(os.path.dirname(OUT), "geo_esearch_esummary_raw.json"), "w") as fh:
        json.dump(raw, fh, indent=1)

    pert = re.compile(r"overexpress|over-express|ectopic|knock-?down|siRNA|shRNA|knock-?out|CRISPR|transfect|inducible|doxycycline|"
                      r"stabl[ey]|silenc|deplet|re-?express|restor", re.I)
    cells = re.compile(r"LNCaP|PC-?3\b|DU-?145|22Rv1|VCaP|RWPE|C4-2|LAPC4|PrEC|organoid|BPH-?1|LHS|MDA-?PCa", re.I)
    rows = []
    for u in uids:
        s = raw["summaries"].get(u, {})
        text = (s.get("title", "") + " " + s.get("summary", ""))
        rows.append({
            "uid": u, "accession": s.get("accession", ""), "title": s.get("title", ""),
            "taxon": s.get("taxon", ""), "gdsType": s.get("gdstype", ""),
            "platform": ";".join("GPL" + p for p in str(s.get("gpl", "")).split(";") if p),
            "n_samples": s.get("n_samples", ""), "pubmed_ids": ";".join(str(p) for p in s.get("pubmedids", [])),
            "release_date": s.get("pdat", ""), "queries": ";".join(sorted(hits[u])),
            "mentions_NKX3_1": bool(re.search(r"NKX3[-.]?1", text, re.I)),
            "perturbation_keyword": ";".join(sorted(set(m.lower() for m in pert.findall(text)))),
            "prostate_cell_model": ";".join(sorted(set(m.upper() for m in cells.findall(text)))),
            "summary": re.sub(r"\s+", " ", s.get("summary", ""))[:1200],
        })
    with open(os.path.join(OUT, "geo_nkx3-1_search_all_hits.tsv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()) if rows else ["uid"], delimiter="\t")
        w.writeheader(); w.writerows(rows)
    print(json.dumps({k: v["count"] for k, v in raw["queries"].items()}, indent=1))
    print("unique series:", len(uids), "accessed", accessed)
    fetch_rwpe1()


# Processed matrices that turned out to be trivially available (DESeq VST values, gene symbols x samples):
#   GSE81439  RWPE1 control vector (CAD1-6) vs NKX3.1 WT (CAD13-18) vs NKX3.1 T164A (CAD91-94)   [Dutta et al. 2016]
#   GSE115335 same design + paraquat: vector (CAD7-12), WT (CAD25-30), R52C (CAD31-36), T164A (CAD95-98)
#             [Papachristodoulou et al. 2021]
SUPPL = {
    "GSE81439": "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE81nnn/GSE81439/suppl/GSE81439_RNAseq_human_RWPE1.txt.gz",
    "GSE115335": "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE115nnn/GSE115335/suppl/GSE115335_RNAseq_human_RWPE1.xlsx",
}
DL = "/mnt/ssd1/Projects/NKX3_wd/work/tcga/geo_download"


def fetch_rwpe1():
    import gzip, hashlib
    os.makedirs(DL, exist_ok=True)
    prov = {}
    for acc, url in SUPPL.items():
        dest = os.path.join(DL, url.rsplit("/", 1)[1])
        if not os.path.exists(dest):
            urllib.request.urlretrieve(url, dest)
            with open(os.path.join(DL, acc + ".accessed_utc"), "w") as fh:
                fh.write(datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"))
        stamp = os.path.join(DL, acc + ".accessed_utc")
        prov[acc] = {"url": url, "local_file": dest, "bytes": os.path.getsize(dest),
                     "md5": hashlib.md5(open(dest, "rb").read()).hexdigest(),
                     "accessed_utc": open(stamp).read().strip() if os.path.exists(stamp) else None}
        if dest.endswith(".xlsx"):           # convert the single sheet to a gzipped TSV for R
            import openpyxl
            tsv = dest.replace(".xlsx", ".tsv.gz")
            if not os.path.exists(tsv):
                wb = openpyxl.load_workbook(dest, read_only=True)
                ws = wb.worksheets[0]
                with gzip.open(tsv, "wt") as out:
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        vals = ["ID_REF" if (i == 0 and j == 0) else ("" if v is None else str(v)) for j, v in enumerate(row)]
                        out.write("\t".join(vals) + "\n")
            prov[acc]["converted_tsv"] = tsv
            prov[acc]["sheet"] = openpyxl.load_workbook(dest, read_only=True).sheetnames
    with open(os.path.join(os.path.dirname(OUT), "geo_rwpe1_download_provenance.json"), "w") as fh:
        json.dump(prov, fh, indent=1)
    print(json.dumps(prov, indent=1))


if __name__ == "__main__":
    main()
